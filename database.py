import sqlite3
import openpyxl
import os


def get_connection():
    return sqlite3.connect('shop.db')


def get_all_products():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT p.id, p.name, p.category, p.description, p.manufacturer, 
               p.supplier, p.price, p.unit, p.quantity, 
               COALESCE(d.percentage, 0) as discount,
               (p.price * (100 - COALESCE(d.percentage, 0)) / 100) as final_price,
               p.image_path
        FROM products p
        LEFT JOIN discounts d ON p.discount_id = d.id
        ORDER BY p.id
    ''')
    result = cursor.fetchall()
    conn.close()
    return result


def get_product_by_id(product_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM products WHERE id = ?', (product_id,))
    result = cursor.fetchone()
    conn.close()
    return result


def add_product(name, category, description, manufacturer, supplier, price, unit, quantity, discount_id, image_path):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO products (name, category, description, manufacturer, supplier, price, unit, quantity, discount_id, image_path)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (name, category, description, manufacturer, supplier, price, unit, quantity, discount_id, image_path))
    conn.commit()
    conn.close()


def update_product(product_id, name, category, description, manufacturer, supplier, price, unit, quantity, discount_id,
                   image_path):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        UPDATE products 
        SET name=?, category=?, description=?, manufacturer=?, supplier=?, 
            price=?, unit=?, quantity=?, discount_id=?, image_path=?
        WHERE id=?
    ''', (
    name, category, description, manufacturer, supplier, price, unit, quantity, discount_id, image_path, product_id))
    conn.commit()
    conn.close()


def delete_product(product_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT COUNT(*) FROM order_items WHERE product_id = ?', (product_id,))
    count = cursor.fetchone()[0]
    if count > 0:
        conn.close()
        return False
    cursor.execute('DELETE FROM products WHERE id = ?', (product_id,))
    conn.commit()
    conn.close()
    return True


def get_all_discounts():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM discounts')
    result = cursor.fetchall()
    conn.close()
    return result


def get_all_orders():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT o.id, o.order_number, o.status, o.delivery_address, o.order_date, o.issue_date, u.full_name
        FROM orders o
        LEFT JOIN users u ON o.user_id = u.id
        ORDER BY o.id
    ''')
    result = cursor.fetchall()
    conn.close()
    return result


def get_order_by_id(order_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM orders WHERE id = ?', (order_id,))
    result = cursor.fetchone()
    conn.close()
    return result


def add_order(order_number, status, delivery_address, order_date, issue_date, user_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO orders (order_number, status, delivery_address, order_date, issue_date, user_id)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (order_number, status, delivery_address, order_date, issue_date, user_id))
    conn.commit()
    conn.close()


def update_order(order_id, order_number, status, delivery_address, order_date, issue_date, user_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        UPDATE orders 
        SET order_number=?, status=?, delivery_address=?, order_date=?, issue_date=?, user_id=?
        WHERE id=?
    ''', (order_number, status, delivery_address, order_date, issue_date, user_id, order_id))
    conn.commit()
    conn.close()


def delete_order(order_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM order_items WHERE order_id = ?', (order_id,))
    cursor.execute('DELETE FROM orders WHERE id = ?', (order_id,))
    conn.commit()
    conn.close()
    return True


def get_unique_manufacturers():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        'SELECT DISTINCT manufacturer FROM products WHERE manufacturer IS NOT NULL AND manufacturer != "" ORDER BY manufacturer')
    result = [row[0] for row in cursor.fetchall()]
    conn.close()
    return result


def get_all_delivery_points():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT id, name, address FROM delivery_points')
    result = cursor.fetchall()
    conn.close()
    return result


# ========== ИМПОРТ ИЗ EXCEL ==========

def import_products_from_excel(file_path):
    """Импорт товаров из Excel файла (Tovar.xlsx)"""
    conn = get_connection()
    cursor = conn.cursor()
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 9:
            continue

        name = row[1] if row[1] else ''
        if not name:
            continue

        category = row[6] if len(row) > 6 and row[6] else ''
        price_str = str(row[3] if len(row) > 3 and row[3] else '0').replace(',', '.').replace('₽', '').strip()
        quantity_str = str(row[8] if len(row) > 8 and row[8] else '0').strip()
        manufacturer = row[5] if len(row) > 5 and row[5] else ''
        supplier = row[4] if len(row) > 4 and row[4] else ''
        discount_str = str(row[7] if len(row) > 7 and row[7] else '0').strip()
        description = row[9] if len(row) > 9 and row[9] else ''
        image = row[10] if len(row) > 10 and row[10] else ''

        try:
            price = float(price_str)
            quantity = int(float(quantity_str))
            discount = int(float(discount_str))
        except:
            price = 0
            quantity = 0
            discount = 0

        discount_id = 1
        if discount > 0:
            cursor.execute('SELECT id FROM discounts WHERE percentage = ?', (discount,))
            disc = cursor.fetchone()
            if disc:
                discount_id = disc[0]
            else:
                cursor.execute('INSERT INTO discounts (name, percentage) VALUES (?, ?)',
                               (f'Скидка {discount}%', discount))
                discount_id = cursor.lastrowid

        try:
            cursor.execute('''
                INSERT INTO products (name, category, price, quantity, description, manufacturer, supplier, unit, discount_id, image_path)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (name, category, price, quantity, description, manufacturer, supplier, 'шт', discount_id, image))
            count += 1
            print(f"✅ Импортирован товар: {name}")
        except Exception as e:
            print(f"❌ Ошибка в строке {row_idx}: {e}")

    conn.commit()
    conn.close()
    print(f"📊 Импортировано товаров: {count}")
    return count


def import_users_from_excel(file_path):
    """Импорт пользователей из Excel файла (user_import.xlsx)"""
    conn = get_connection()
    cursor = conn.cursor()
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 4:
            continue

        role_name = str(row[0]).strip().lower() if row[0] else ''
        full_name = str(row[1]).strip() if row[1] else ''
        login = str(row[2]).strip() if row[2] else ''
        password = str(row[3]).strip() if row[3] else ''

        if 'администратор' in role_name:
            role = 'admin'
        elif 'менеджер' in role_name:
            role = 'manager'
        else:
            role = 'client'

        if not login:
            continue

        try:
            cursor.execute('INSERT OR REPLACE INTO users (login, password, full_name, role) VALUES (?, ?, ?, ?)',
                           (login, password, full_name, role))
            count += 1
            print(f"✅ Импортирован пользователь: {login} ({role})")
        except Exception as e:
            print(f"❌ Ошибка в строке {row_idx}: {e}")

    conn.commit()
    conn.close()
    print(f"📊 Импортировано пользователей: {count}")
    return count


def import_delivery_points_from_excel(file_path):
    """Импорт пунктов выдачи из Excel файла (Пункты выдачи_import.xlsx)"""
    conn = get_connection()
    cursor = conn.cursor()
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if not row or not row[0]:
            continue

        address = str(row[0]).strip()
        if not address:
            continue

        name = f"Пункт {row_idx}"

        try:
            cursor.execute('INSERT INTO delivery_points (name, address) VALUES (?, ?)', (name, address))
            count += 1
            print(f"✅ Импортирован пункт выдачи: {name} - {address}")
        except Exception as e:
            print(f"❌ Ошибка в строке {row_idx}: {e}")

    conn.commit()
    conn.close()
    print(f"📊 Импортировано пунктов выдачи: {count}")
    return count


def import_orders_from_excel(file_path):
    """Импорт заказов из Excel файла (Заказ_import.xlsx)"""
    conn = get_connection()
    cursor = conn.cursor()
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or len(row) < 7:
            continue

        order_num = row[0] if row[0] else ''
        article_str = str(row[1]) if row[1] else ''
        order_date = row[2] if row[2] else ''
        issue_date = row[3] if row[3] else ''
        delivery_address = row[4] if row[4] else ''
        client_name = row[5] if row[5] else ''
        code = row[6] if row[6] else ''

        if not order_num:
            continue

        # Находим user_id по ФИО
        cursor.execute('SELECT id FROM users WHERE full_name = ?', (client_name,))
        user = cursor.fetchone()
        user_id = user[0] if user else 3

        try:
            cursor.execute('''
                INSERT INTO orders (order_number, status, delivery_address, order_date, issue_date, user_id)
                VALUES (?, ?, ?, ?, ?, ?)
            ''', (f"ORDER-{order_num}", 'Новый', delivery_address, order_date, issue_date, user_id))
            count += 1
            print(f"✅ Импортирован заказ: {order_num}")
        except Exception as e:
            print(f"❌ Ошибка в строке {row_idx}: {e}")

    conn.commit()
    conn.close()
    print(f"📊 Импортировано заказов: {count}")
    return count